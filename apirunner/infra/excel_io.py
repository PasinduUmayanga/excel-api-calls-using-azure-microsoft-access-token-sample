"""Read and write Excel/CSV files with openpyxl."""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field

from apirunner.core.models import CallStatus, FieldSpec, Job
from apirunner.core.validation import RowValidation

HEADER_FILL = PatternFill("solid", fgColor="6D28D9")
SUCCESS_FILL = PatternFill("solid", fgColor="DCFCE7")
FAILED_FILL = PatternFill("solid", fgColor="FEE2E2")
ERROR_FILL = PatternFill("solid", fgColor="FCA5A5")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RESULT_COLUMNS = ["_status", "_http_code", "_error", "_response", "_attempts"]
MAX_CELL = 32000  # Excel limit is 32767 chars


class Table(BaseModel):
    sheets: list[str] = Field(default_factory=list)
    sheet: str = ""
    columns: list[str] = Field(default_factory=list)
    rows: list[dict[str, Any]] = Field(default_factory=list)


def read_table(content: bytes, filename: str, sheet: str | None = None) -> Table:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        columns = [c for c in (reader.fieldnames or []) if c]
        rows = [{c: (r.get(c) if r.get(c) != "" else None) for c in columns} for r in reader]
        return Table(sheets=["csv"], sheet="csv", columns=columns, rows=_drop_empty(rows))

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None) or ()
    columns = [str(h).strip() if h is not None else "" for h in header]
    rows = []
    for values in it:
        rows.append({c: v for c, v in zip(columns, values) if c})
    named = [c for c in columns if c]
    table = Table(sheets=list(wb.sheetnames), sheet=ws.title, columns=named, rows=_drop_empty(rows))
    wb.close()
    return table


def _drop_empty(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [r for r in rows if any(v not in (None, "") for v in r.values())]


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (int, float, bool, date, datetime)):
        return value
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    return str(value)[:MAX_CELL]


def _style_header(ws, count: int) -> None:
    for i in range(1, count + 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = 22
    ws.freeze_panes = "A2"


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_template(fields: list[FieldSpec], sheet_name: str = "Data") -> bytes:
    """Excel template whose header row is exactly the field names, with type notes and dropdowns."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append([f.name for f in fields])
    _style_header(ws, len(fields))
    for i, f in enumerate(fields, start=1):
        note = f"type: {f.type}\n{'required' if f.required else 'optional'}"
        if f.regex:
            note += f"\npattern: {f.regex}"
        if f.allowed:
            note += f"\nallowed: {', '.join(f.allowed)}"
            dv = DataValidation(type="list", formula1='"' + ",".join(f.allowed) + '"', allow_blank=not f.required)
            ws.add_data_validation(dv)
            col = get_column_letter(i)
            dv.add(f"{col}2:{col}5000")
        ws.cell(row=1, column=i).comment = Comment(note, "API Runner")
    return _save(wb)


def write_results(job: Job, status: CallStatus) -> bytes:
    """Success or Failed workbook: original columns + result columns. Re-uploadable as-is."""
    wb = Workbook()
    ws = wb.active
    ws.title = status.value.capitalize()
    columns = [c for c in job.columns if not c.startswith("_")]
    ws.append(columns + RESULT_COLUMNS)
    _style_header(ws, len(columns) + len(RESULT_COLUMNS))
    fill = SUCCESS_FILL if status == CallStatus.SUCCESS else FAILED_FILL
    for row in job.rows:
        if row.result.status != status:
            continue
        r = row.result
        ws.append(
            [_cell(row.source.get(c)) for c in columns]
            + [r.status.value, r.http_code, r.error, _cell(r.response_text), r.attempts]
        )
        for cell in ws[ws.max_row][len(columns):]:
            cell.fill = fill
    return _save(wb)


def write_invalid_rows(columns: list[str], rows: list[dict[str, Any]], mapping: dict[str, str | None],
                       validations: list[RowValidation]) -> bytes:
    """Invalid rows with the bad cells highlighted and an _error column, ready to fix and re-upload."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invalid"
    cols = [c for c in columns if not c.startswith("_")]
    ws.append(cols + ["_error"])
    _style_header(ws, len(cols) + 1)
    for v in validations:
        if v.valid:
            continue
        source = rows[v.index]
        ws.append([_cell(source.get(c)) for c in cols] + ["; ".join(f"{k}: {m}" for k, m in v.errors.items())])
        for field_name in v.errors:
            col = mapping.get(field_name)
            if col in cols:
                ws.cell(row=ws.max_row, column=cols.index(col) + 1).fill = ERROR_FILL
    return _save(wb)
