import io

from openpyxl import Workbook, load_workbook

from apirunner.core.mapping import auto_map
from apirunner.core.models import CallResult, CallStatus, FieldSpec, RequestTemplate
from apirunner.core.executor import create_job
from apirunner.core.validation import validate_rows
from apirunner.infra.excel_io import RESULT_COLUMNS, read_table, write_invalid_rows, write_results, write_template

FIELDS = [FieldSpec(name="name"), FieldSpec(name="age", type="integer"),
          FieldSpec(name="plan", allowed=["free", "pro"], required=False)]


def xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_xlsx_skips_blank_rows_and_unnamed_columns():
    table = read_table(xlsx([["name", "age", None], ["Ann", 31, "x"], [None, None, None], ["Bob", 4, None]]),
                       "f.xlsx")
    assert table.columns == ["name", "age"]
    assert table.rows == [{"name": "Ann", "age": 31}, {"name": "Bob", "age": 4}]


def test_read_csv_with_bom():
    table = read_table("﻿name,age\nAnn,31\n,\n".encode(), "f.csv")
    assert table.columns == ["name", "age"] and table.rows == [{"name": "Ann", "age": "31"}]


def test_template_has_exact_headers_notes_and_dropdowns():
    wb = load_workbook(io.BytesIO(write_template(FIELDS)))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["name", "age", "plan"]
    assert "integer" in ws["B1"].comment.text
    assert ws.data_validations.dataValidation[0].formula1 == '"free,pro"'


def test_template_reads_back_and_maps_one_to_one():
    table = read_table(write_template(FIELDS), "t.xlsx")
    assert auto_map(table.columns, FIELDS) == {"name": "name", "age": "age", "plan": "plan"}


def _job():
    rows = [{"name": "Ann", "age": 1}, {"name": "Bob", "age": 2}]
    template = RequestTemplate(fields=FIELDS[:2])
    mapping = {"name": "name", "age": "age"}
    job = create_job(template, ["name", "age"], mapping, rows, validate_rows(rows, mapping, template.fields))
    job.rows[0].result = CallResult(status=CallStatus.SUCCESS, http_code=201, response_text='{"id": 1}', attempts=1)
    job.rows[1].result = CallResult(status=CallStatus.FAILED, http_code=500, error="HTTP 500", attempts=3)
    return job


def test_results_split_into_success_and_failed():
    job = _job()
    ok = read_table(write_results(job, CallStatus.SUCCESS), "Success.xlsx")
    failed = read_table(write_results(job, CallStatus.FAILED), "Failed.xlsx")
    assert ok.columns == ["name", "age"] + RESULT_COLUMNS
    assert [r["name"] for r in ok.rows] == ["Ann"] and [r["name"] for r in failed.rows] == ["Bob"]
    assert failed.rows[0]["_http_code"] == 500 and failed.rows[0]["_attempts"] == 3


def test_failed_file_can_be_re_uploaded_and_validates():
    failed = read_table(write_results(_job(), CallStatus.FAILED), "Failed.xlsx")
    mapping = auto_map(failed.columns, FIELDS[:2])
    assert mapping == {"name": "name", "age": "age"}  # _status etc. are ignored
    assert all(v.valid for v in validate_rows(failed.rows, mapping, FIELDS[:2]))


def test_invalid_rows_export_has_reasons():
    rows = [{"name": "Ann", "age": "x"}, {"name": "Bob", "age": 2}]
    mapping = {"name": "name", "age": "age"}
    vals = validate_rows(rows, mapping, FIELDS[:2])
    out = read_table(write_invalid_rows(["name", "age"], rows, mapping, vals), "Invalid.xlsx")
    assert out.rows == [{"name": "Ann", "age": "x", "_error": "age: expected integer, got 'x'"}]
